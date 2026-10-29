"""Read supported files into plain-text pages with source metadata.

The RAG pipeline only needs text, but documents arrive in a number of common
formats.  Each format-specific dependency is imported when that format is
used so that the text/CSV/JSON paths stay lightweight and failures can explain
which optional package is missing.
"""

from __future__ import annotations

import csv
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


# Keep this set public: the CLI imports it when recursively indexing a folder.
TEXT_EXTENSIONS = {
    ".asm",
    ".bash",
    ".bat",
    ".c",
    ".cfg",
    ".cmd",
    ".conf",
    ".cpp",
    ".cs",
    ".css",
    ".go",
    ".h",
    ".hpp",
    ".ini",
    ".java",
    ".js",
    ".jsx",
    ".jsonl",
    ".kts",
    ".kt",
    ".less",
    ".log",
    ".md",
    ".mdx",
    ".php",
    ".pl",
    ".properties",
    ".ps1",
    ".py",
    ".rb",
    ".rest",
    ".rst",
    ".rs",
    ".sass",
    ".scss",
    ".sh",
    ".sql",
    ".svg",
    ".swift",
    ".tex",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}

SUPPORTED = {
    ".csv",
    ".docx",
    ".html",
    ".htm",
    ".json",
    ".pdf",
    ".pptx",
    ".tsv",
    ".xls",
    ".xlsx",
    ".xlsm",
} | TEXT_EXTENSIONS


class LoaderDependencyError(RuntimeError):
    """Raised when a supported format's optional reader is not installed."""


def _missing_dependency(extension: str, package: str, install_name: str | None = None) -> LoaderDependencyError:
    name = install_name or package
    return LoaderDependencyError(
        f"Reading {extension} files requires the optional '{package}' package. "
        f"Install it with: pip install {name}"
    )


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="ignore").strip()


_OCR_PROMPT = (
    "You are an OCR engine. Extract ALL text visible in this document exactly "
    "as written. Include names, numbers, dates, labels, and table cells. "
    "Preserve reading order and line breaks. Do not summarize, translate, "
    "explain, or add commentary. If a value is unclear, give your best "
    "transcription. Return plain text only, without Markdown fences."
)


def _ocr_pdf_with_gemini(path: Path, source: str) -> list[dict[str, Any]]:
    """OCR a scanned PDF by sending it to Gemini as an inline document.

    Returns an empty list when OCR is disabled or no Gemini key is configured;
    the caller then keeps whatever native text pypdf managed to find.
    """
    import base64
    import os

    if os.getenv("OCR_ENABLED", "true").strip().lower() in {"0", "false", "no", "off"}:
        return []
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        return []

    import requests

    size = path.stat().st_size
    if size > 18 * 1024 * 1024:
        raise ValueError(
            f"'{source}' looks scanned and is {size / 1_000_000:.1f} MB, which is "
            "too large for inline OCR. Upload a PDF under 18 MB or add a text layer."
        )
    model = (
        os.getenv("GEMINI_OCR_MODEL", "").strip()
        or os.getenv("LLM_MODEL", "").strip()
        or "gemini-2.5-flash"
    )
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    body = {
        "contents": [
            {
                "role": "user",
                "parts": [
                    {"text": _OCR_PROMPT},
                    {
                        "inline_data": {
                            "mime_type": "application/pdf",
                            "data": base64.b64encode(path.read_bytes()).decode("ascii"),
                        }
                    },
                ],
            }
        ],
        "generationConfig": {"temperature": 0.0, "maxOutputTokens": 8192},
    }
    try:
        resp = requests.post(url, params={"key": api_key}, json=body, timeout=180)
    except requests.exceptions.RequestException as exc:
        raise ValueError(
            f"OCR for '{source}' failed: could not reach the Gemini API."
        ) from exc
    if resp.status_code >= 400:
        detail = resp.text[:300]
        try:
            detail = resp.json().get("error", {}).get("message", detail)
        except Exception:
            pass
        raise ValueError(
            f"OCR for '{source}' failed (HTTP {resp.status_code}): {detail}"
        )
    try:
        candidates = resp.json().get("candidates") or []
        parts = candidates[0].get("content", {}).get("parts", []) if candidates else []
        text = "\n".join(str(part.get("text", "")) for part in parts).strip()
    except (TypeError, ValueError, KeyError) as exc:
        raise ValueError(f"OCR for '{source}' returned an unexpected response.") from exc
    text = re.sub(r"^```(?:text)?\s*|\s*```$", "", text).strip()
    return [{"text": text, "source": source, "page": 1}] if text else []


def _load_pdf(path: Path, source: str) -> list[dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise _missing_dependency(".pdf", "pypdf") from exc

    pages: list[dict[str, Any]] = []
    page_count = 0
    try:
        with path.open("rb") as handle:
            reader = PdfReader(handle)
            page_count = len(reader.pages)
            for i, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append({"text": text, "source": source, "page": i})
    except Exception as exc:
        raise ValueError(f"Could not extract text from PDF '{source}': {exc}") from exc

    # Image-only scans often expose just a header or a page number. OCR when
    # the native text layer is too thin to answer questions about the document.
    extracted_chars = sum(len(page["text"]) for page in pages)
    if page_count and extracted_chars < max(40, 20 * page_count):
        ocr_pages = _ocr_pdf_with_gemini(path, source)
        if sum(len(page["text"]) for page in ocr_pages) > extracted_chars:
            return ocr_pages
    return pages


def _load_docx(path: Path, source: str) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise _missing_dependency(".docx", "python-docx", "python-docx") from exc

    try:
        document = Document(str(path))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append("\t".join(values))
    except Exception as exc:
        raise ValueError(f"Could not extract text from DOCX '{source}': {exc}") from exc

    text = "\n".join(parts).strip()
    return [{"text": text, "source": source, "page": None}] if text else []


def _shape_text(shape: Any) -> list[str]:
    """Extract text from normal PowerPoint shapes and table cells."""
    values: list[str] = []
    text = getattr(shape, "text", None)
    if text and text.strip():
        values.append(text.strip())
    table = getattr(shape, "table", None)
    if table is not None:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                values.append("\t".join(cells))
    return values


def _load_pptx(path: Path, source: str) -> list[dict[str, Any]]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise _missing_dependency(".pptx", "python-pptx", "python-pptx") from exc

    pages: list[dict[str, Any]] = []
    try:
        presentation = Presentation(str(path))
        for i, slide in enumerate(presentation.slides, start=1):
            values: list[str] = []
            for shape in slide.shapes:
                values.extend(_shape_text(shape))
            text = "\n".join(values).strip()
            if text:
                pages.append({"text": text, "source": source, "page": i})
    except Exception as exc:
        raise ValueError(f"Could not extract text from PPTX '{source}': {exc}") from exc
    return pages


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _load_xlsx(path: Path, source: str) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise _missing_dependency(".xlsx", "openpyxl") from exc

    pages: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows: list[str] = []
                for row in worksheet.iter_rows(values_only=True):
                    values = [_cell_text(value) for value in row]
                    while values and not values[-1]:
                        values.pop()
                    if any(values):
                        rows.append("\t".join(values))
                if rows:
                    pages.append(
                        {
                            "text": f"Sheet: {worksheet.title}\n" + "\n".join(rows),
                            "source": source,
                            # A sheet name is more useful to a citation than a
                            # made-up page number and is supported by Chroma's
                            # string metadata values.
                            "page": worksheet.title,
                            "section": worksheet.title,
                        }
                    )
        finally:
            close = getattr(workbook, "close", None)
            if callable(close):
                close()
    except Exception as exc:
        raise ValueError(f"Could not extract text from XLSX '{source}': {exc}") from exc
    return pages


def _load_xls(path: Path, source: str) -> list[dict[str, Any]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise _missing_dependency(".xls", "xlrd") from exc

    pages: list[dict[str, Any]] = []
    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for worksheet in workbook.sheets():
                rows: list[str] = []
                for row_index in range(worksheet.nrows):
                    values = [_cell_text(value) for value in worksheet.row_values(row_index)]
                    while values and not values[-1]:
                        values.pop()
                    if any(values):
                        rows.append("\t".join(values))
                if rows:
                    pages.append(
                        {
                            "text": f"Sheet: {worksheet.name}\n" + "\n".join(rows),
                            "source": source,
                            "page": worksheet.name,
                            "section": worksheet.name,
                        }
                    )
        finally:
            release = getattr(workbook, "release_resources", None)
            if callable(release):
                release()
    except Exception as exc:
        raise ValueError(f"Could not extract text from XLS '{source}': {exc}") from exc
    return pages


def _load_delimited(path: Path, source: str, extension: str) -> list[dict[str, Any]]:
    text = _read_text(path)
    if not text:
        return []

    delimiter = "\t" if extension == ".tsv" else ","
    if extension == ".csv":
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            pass

    try:
        rows = csv.reader(io.StringIO(text), delimiter=delimiter)
        formatted = ["\t".join(cell.strip() for cell in row).rstrip() for row in rows]
    except csv.Error as exc:
        raise ValueError(f"Could not parse {extension} file '{source}': {exc}") from exc

    output = "\n".join(row for row in formatted if row).strip()
    return [{"text": output, "source": source, "page": None}] if output else []


class _HTMLTextExtractor(HTMLParser):
    """Small stdlib HTML-to-text reader; no BeautifulSoup dependency needed."""

    _BLOCK_TAGS = {
        "address",
        "article",
        "aside",
        "blockquote",
        "br",
        "div",
        "dl",
        "dt",
        "dd",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "table",
        "td",
        "th",
        "tr",
        "ul",
    }
    _IGNORED_TAGS = {"script", "style", "noscript", "template"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in self._IGNORED_TAGS:
            self._ignored_depth += 1
        if not self._ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self._IGNORED_TAGS and self._ignored_depth:
            self._ignored_depth -= 1
        if not self._ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)

    def text(self) -> str:
        value = "".join(self.parts)
        value = re.sub(r"[ \t\r\f\v]+", " ", value)
        value = re.sub(r"\n[ ]+", "\n", value)
        return re.sub(r"\n{3,}", "\n\n", value).strip()


def _load_html(path: Path, source: str) -> list[dict[str, Any]]:
    parser = _HTMLTextExtractor()
    try:
        parser.feed(_read_text(path))
        parser.close()
    except Exception as exc:
        raise ValueError(f"Could not extract text from HTML '{source}': {exc}") from exc
    text = parser.text()
    return [{"text": text, "source": source, "page": None}] if text else []


def _load_json(path: Path, source: str) -> list[dict[str, Any]]:
    raw = _read_text(path)
    if not raw:
        return []
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Could not parse JSON '{source}': {exc.msg} at line {exc.lineno}") from exc
    text = json.dumps(value, ensure_ascii=False, indent=2)
    return [{"text": text, "source": source, "page": None}] if text else []


def load_file(path: str | Path, source: str | None = None) -> list[dict[str, Any]]:
    """Return ``[{"text", "source", "page"}]`` pages for *path*.

    ``page`` is one-based for PDF/PPTX pages, a worksheet name for XLSX, and
    ``None`` for formats without a useful page concept.  ``source`` can be
    supplied by callers that load a temporary file but need citations to use
    the uploaded filename.
    """

    path = Path(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {extension or '(none)'}. Use {sorted(SUPPORTED)}")
    source_name = source or path.name

    if extension == ".pdf":
        return _load_pdf(path, source_name)
    if extension in TEXT_EXTENSIONS:
        text = _read_text(path)
        return [{"text": text, "source": source_name, "page": None}] if text else []
    if extension == ".docx":
        return _load_docx(path, source_name)
    if extension == ".pptx":
        return _load_pptx(path, source_name)
    if extension in {".xlsx", ".xlsm"}:
        return _load_xlsx(path, source_name)
    if extension == ".xls":
        return _load_xls(path, source_name)
    if extension in {".csv", ".tsv"}:
        return _load_delimited(path, source_name, extension)
    if extension in {".html", ".htm"}:
        return _load_html(path, source_name)
    if extension == ".json":
        return _load_json(path, source_name)

    # The extension check above makes this unreachable, but retaining a clear
    # error protects callers if SUPPORTED is changed without a loader branch.
    raise ValueError(f"Unsupported file type: {extension}")
